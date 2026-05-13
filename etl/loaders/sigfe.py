"""Loader SIGFE 2.0 — Detecta y carga reportes Jasper (.xls) y nativos (.xlsx).

Maneja:
- Conversión automática .xls binario Jasper → .xlsx (vía libreoffice)
- Detección automática de fila de headers (varía entre reportes)
- Parsing de metadata (período, fecha extracción, cobertura)
- Normalización a DataFrame consistente

Uso típico:
    from etl.loaders.sigfe import load_sigfe_report
    df, meta = load_sigfe_report("path/to/SB_ListadoPagosRealizados__1_.xls")
"""
from pathlib import Path
import subprocess
import tempfile
import warnings
from typing import Tuple, Optional

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

# Palabras clave que identifican una fila de headers en reportes SIGFE Jasper
HEADER_HINTS = {
    'Área Transaccional', 'Unidad Ejecutora', 'Código Unidad Ejecutora',
    'Folio', 'Cuenta Contable', 'Tipo Operación', 'Tipo Presupuesto',
    'Folio Requerimiento', 'Catálogo 01', 'Concepto Presupuesto',
    'Concepto Presupuestario', 'Nivel', 'Fecha', 'ID',
    'Cuenta', 'Concepto', 'Principal', 'Tipo Factura'
}


def _convert_xls_to_xlsx(xls_path: Path, out_dir: Path) -> Path:
    """Convierte Jasper .xls a .xlsx vía libreoffice headless."""
    result = subprocess.run(
        ['libreoffice', '--headless', '--convert-to', 'xlsx',
         '--outdir', str(out_dir), str(xls_path)],
        capture_output=True, text=True, timeout=120
    )
    if result.returncode != 0:
        raise RuntimeError(f"Falló conversión: {result.stderr}")
    xlsx_path = out_dir / (xls_path.stem + '.xlsx')
    if not xlsx_path.exists():
        raise FileNotFoundError(f"No se creó archivo: {xlsx_path}")
    return xlsx_path


def _is_header_row(row) -> bool:
    """Identifica si una fila es de headers."""
    cells = [str(c).strip() if c is not None else '' for c in row]
    non_empty = [c for c in cells if c]
    if len(non_empty) < 3:
        return False
    return any(h in cells[:6] for h in HEADER_HINTS)


def _ensure_xlsx(path: Path) -> Tuple[Path, Optional[Path]]:
    """Asegura tener .xlsx. Devuelve (path_xlsx, path_temp_si_creado).
    El segundo elemento se usa para cleanup."""
    if path.suffix.lower() == '.xlsx':
        return path, None
    elif path.suffix.lower() == '.xls':
        tmpdir = Path(tempfile.mkdtemp(prefix='sigfe_'))
        xlsx = _convert_xls_to_xlsx(path, tmpdir)
        return xlsx, tmpdir
    raise ValueError(f"Extensión no soportada: {path.suffix}")


def load_sigfe_report(path: str | Path, sheet: Optional[str] = None) -> Tuple[pd.DataFrame, dict]:
    """Carga un reporte SIGFE como DataFrame.

    Args:
        path: ruta al archivo .xls o .xlsx
        sheet: nombre de hoja (si None, usa la primera)

    Returns:
        (df, metadata): DataFrame con datos + dict con metadata
    """
    path = Path(path)
    xlsx_path, tmpdir = _ensure_xlsx(path)

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

        # Cargar todas las filas; algunas hojas tienen contenido fuera del rango
        # declarado por max_row, por lo que iter_rows() de read_only puede fallar
        rows = []
        max_r = 0; max_c = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    if cell.row > max_r: max_r = cell.row
                    if cell.column > max_c: max_c = cell.column
        if max_r == 0:
            return pd.DataFrame(), {'archivo': str(path), 'vacio': True}

        for r in range(1, max_r + 1):
            row = tuple(ws.cell(row=r, column=c).value for c in range(1, max_c + 1))
            rows.append(row)

        # Detectar fila de headers
        header_idx = None
        for i, row in enumerate(rows):
            if _is_header_row(row):
                header_idx = i
                break

        if header_idx is None:
            # Reporte no estándar — devolver todo como DataFrame raw
            return pd.DataFrame(rows), {
                'archivo': str(path), 'hoja': ws.title, 'header_detected': False
            }

        # Metadata: filas previas a headers
        metadata = {
            'archivo': str(path),
            'hoja': ws.title,
            'titulo': rows[0][0] if rows else None,
            'institucion': rows[1][0] if len(rows) > 1 else None,
            'contexto': [r[0] for r in rows[2:header_idx] if r and r[0] is not None],
            'header_row': header_idx + 1,
        }

        # Construir headers
        headers_raw = list(rows[header_idx])
        headers = []
        for i, h in enumerate(headers_raw):
            if h is None or str(h).strip() == '':
                headers.append(f'col_{i}')
            else:
                headers.append(str(h).strip())
        # Recortar columnas vacías al final
        while headers and headers[-1].startswith('col_'):
            headers.pop()
        n_cols = len(headers)

        # Data
        data_rows = []
        for r in rows[header_idx + 1:]:
            if r and r[0] is not None and str(r[0]).strip() not in ('', 'Total', 'TOTAL'):
                data_rows.append(r[:n_cols])

        df = pd.DataFrame(data_rows, columns=headers)
        metadata['n_filas'] = len(df)
        metadata['n_cols'] = len(df.columns)

        return df, metadata

    finally:
        if tmpdir is not None:
            import shutil
            shutil.rmtree(tmpdir, ignore_errors=True)


def load_sigfe_native_xlsx(path: str | Path, sheet: str = 'Sheet1') -> Tuple[pd.DataFrame, dict]:
    """Carga reportes SIGFE en formato xlsx nativo.
    
    Estos vienen con estructuras variadas:
    - Pago_de_Facturas_Chile_Paga: headers en F1
    - Estado_de_Situación_Presupuestaria: headers en F1 (algunos vacíos)
    - Estado_Compromiso: headers en F1
    - Tesorería: estructura matricial (no estándar)
    - Cartera_Financiera_*: headers en F1
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]

    rows = []
    max_r = 0; max_c = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                if cell.row > max_r: max_r = cell.row
                if cell.column > max_c: max_c = cell.column

    if max_r == 0:
        return pd.DataFrame(), {'archivo': str(path), 'vacio': True}

    for r in range(1, max_r + 1):
        rows.append(tuple(ws.cell(row=r, column=c).value for c in range(1, max_c + 1)))

    # Detectar headers: primera fila con >= 2 celdas no vacías
    header_idx = None
    for i, row in enumerate(rows):
        non_empty = [c for c in row if c is not None and str(c).strip()]
        if len(non_empty) >= 2:
            header_idx = i
            break

    if header_idx is None:
        return pd.DataFrame(), {'archivo': str(path), 'vacio': True}

    headers_raw = list(rows[header_idx])
    headers = [str(c).strip() if c else f'col_{i}' for i, c in enumerate(headers_raw)]
    while headers and headers[-1].startswith('col_'):
        headers.pop()
    n_cols = len(headers)

    data = [r[:n_cols] for r in rows[header_idx + 1:]
            if any(c is not None and str(c).strip() for c in r[:n_cols])]
    df = pd.DataFrame(data, columns=headers)

    metadata = {
        'archivo': str(path),
        'hoja': sheet,
        'titulo': path.stem,
        'header_row': header_idx + 1,
        'n_filas': len(df),
        'n_cols': len(df.columns),
    }

    if 'Sheet2' in wb.sheetnames and sheet != 'Sheet2':
        ws_meta = wb['Sheet2']
        meta_pairs = {}
        for row in ws_meta.iter_rows(values_only=True):
            if row and row[0] and len(row) > 1 and row[1]:
                meta_pairs[str(row[0]).strip()] = str(row[1]).strip()
        metadata['filtros'] = meta_pairs

    return df, metadata


if __name__ == '__main__':
    # Test rápido
    import sys
    if len(sys.argv) > 1:
        for p in sys.argv[1:]:
            path = Path(p)
            if path.suffix == '.xls':
                df, meta = load_sigfe_report(path)
            else:
                df, meta = load_sigfe_native_xlsx(path)
            print(f"\n📄 {path.name}")
            print(f"   Título: {meta.get('titulo')}")
            print(f"   Forma: {df.shape}")
            print(f"   Cols: {list(df.columns)[:8]}...")
